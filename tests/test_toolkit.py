"""
Tests for the DPDPA audit toolkit.

    python -m pytest -q            (or: python -m unittest discover tests)

The suite is written against the properties that matter in an audit rather than
against implementation detail: that the library validates and refuses to load
when it does not, that scoping is declarative and reproducible, that "not
applicable" leaves the denominator rather than scoring zero, that the weighted
aggregate is genuinely domain-weighted, that risk and priority follow the stated
bands, that penalty figures are ceilings drawn from the Schedule, and that no
unverified Rules citation can escape as a plausible-looking rule number.
"""

from __future__ import annotations

import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dpdpa_audit import (  # noqa: E402
    PROFILE_FLAGS,
    SUBSTANTIALLY_COMPLIANT,
    Assessment,
    CompanyProfile,
    build_gap_register,
    load_library,
    maturity_band,
    priority_for,
    risk_band,
    score_assessment,
)
from dpdpa_audit.core import NOT_APPLICABLE  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTROLS = os.path.join(ROOT, "controls")
SAMPLE = os.path.join(ROOT, "samples", "sample_assessment.json")

LIB = load_library(CONTROLS)


def _profile(**flags) -> CompanyProfile:
    p = CompanyProfile.blank()
    p.legal_name = "Test Fiduciary Private Limited"
    p.flags.update(flags)
    return p


class TestLibraryIntegrity(unittest.TestCase):
    def test_loads_and_is_populated(self):
        self.assertEqual(len(LIB.domains), 17)
        self.assertEqual(len(LIB.controls), 90)

    def test_every_control_belongs_to_a_declared_domain(self):
        for c in LIB.controls.values():
            self.assertIn(c.domain, LIB.domains, f"{c.id} has unknown domain {c.domain}")

    def test_control_ids_are_prefixed_with_their_domain(self):
        for c in LIB.controls.values():
            self.assertTrue(c.id.startswith(c.domain + "-"), c.id)

    def test_domain_control_counts_sum_to_the_library(self):
        total = sum(len(LIB.controls_in(d.code)) for d in LIB.ordered_domains())
        self.assertEqual(total, len(LIB.controls))

    def test_every_control_cites_the_act(self):
        for c in LIB.controls.values():
            self.assertTrue(c.act_ref.strip(), f"{c.id} has no Act reference")
            self.assertIn("s.", c.act_ref, f"{c.id} Act reference looks malformed")

    def test_scoring_anchors_are_present_and_described(self):
        for c in LIB.controls.values():
            for level in ("0", "2", "4"):
                self.assertIn(level, c.scoring, f"{c.id} lacks anchor {level}")
                self.assertTrue(c.scoring[level].strip(), f"{c.id} anchor {level} is empty")

    def test_target_state_is_the_level_four_anchor(self):
        for c in LIB.controls.values():
            self.assertEqual(c.target_state, c.scoring["4"])

    def test_severities_are_from_the_closed_set(self):
        for c in LIB.controls.values():
            self.assertIn(c.severity, ("critical", "high", "medium", "low"), c.id)

    def test_applicability_flags_are_declared_in_profile_flags(self):
        for c in LIB.controls.values():
            for key in ("requires", "excluded_if"):
                for flag in c.applicability.get(key, []):
                    self.assertIn(flag, PROFILE_FLAGS, f"{c.id} uses unknown flag {flag}")

    def test_rule_slots_referenced_by_controls_exist(self):
        for c in LIB.controls.values():
            if c.rule_slot:
                self.assertIn(c.rule_slot, LIB.citations["rule_slots"], c.id)

    def test_penalty_entries_referenced_by_controls_exist(self):
        entries = {e["entry"] for e in LIB.citations["penalty_schedule"]["entries"]}
        for c in LIB.controls.values():
            if c.penalty_entry is not None:
                self.assertIn(c.penalty_entry, entries, c.id)

    def test_a_bad_library_raises_rather_than_loading_quietly(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in os.listdir(CONTROLS):
                src = os.path.join(CONTROLS, name)
                with open(src, encoding="utf-8") as fh:
                    body = fh.read()
                if name == "01_foundation.yaml":
                    body = body.replace("severity: critical", "severity: catastrophic", 1)
                with open(os.path.join(tmp, name), "w", encoding="utf-8") as fh:
                    fh.write(body)
            with self.assertRaises(ValueError) as ctx:
                load_library(tmp)
            self.assertIn("catastrophic", str(ctx.exception))


class TestCitationPolicy(unittest.TestCase):
    """No unverified Rules citation may leave the toolkit looking authoritative."""

    PLACEHOLDER = "[Rules citation to be verified]"

    def test_unverified_slots_render_as_a_placeholder(self):
        for c in LIB.controls.values():
            if c.rule_slot and LIB.rule_status(c) == "VERIFY_AGAINST_NOTIFIED_RULES":
                self.assertEqual(LIB.rule_citation_text(c), self.PLACEHOLDER, c.id)

    def test_no_slot_carries_a_rule_number_while_unverified(self):
        for key, slot in LIB.citations["rule_slots"].items():
            if slot.get("status") == "VERIFY_AGAINST_NOTIFIED_RULES":
                self.assertIn(slot.get("citation") or "", ("", None), key)

    def test_every_unverified_slot_states_its_subject_and_act_anchor(self):
        for s in LIB.unverified_slots():
            raw = LIB.citations["rule_slots"][s["key"]]
            self.assertTrue(s["subject"].strip(), s["key"])
            self.assertTrue(str(raw.get("anchors_act") or "").strip(), s["key"])

    def test_unverified_slots_list_the_controls_relying_on_them(self):
        listed = {c for s in LIB.unverified_slots() for c in s["controls"]}
        expected = {c.id for c in LIB.controls.values()
                    if c.rule_slot and LIB.rule_status(c) == "VERIFY_AGAINST_NOTIFIED_RULES"}
        self.assertEqual(listed, expected)

    def test_penalty_schedule_is_marked_verified_against_the_act(self):
        self.assertEqual(LIB.citations["penalty_schedule"]["status"], "VERIFIED_ACT")

    def test_no_rules_material_leaks_into_the_library(self):
        """The Act is cited; the Rules are not, anywhere, in any field.

        The slots keep unverified rule numbers out of the citation fields, but a
        prescribed threshold can just as easily arrive in prose — a notification
        window in an audit question, a net-worth figure in a scoring anchor. Those
        would read as verified law because nothing marks them as unverified. So
        the whole library is scanned, not just the citation registry.
        """
        parts = []
        for f in sorted(os.listdir(CONTROLS)):
            if f.endswith(".yaml"):
                with open(os.path.join(CONTROLS, f), encoding="utf-8") as fh:
                    parts.append(fh.read())
        blob = "\n".join(parts)
        forbidden = [
            (r"\brule[s]?\s+\d", "a rule number"),
            (r"\b\d+\s*(?:hours|hrs)\b", "a prescribed notification window"),
            (r"\bwithin\s+\d+\s+(?:days|months)\b", "a prescribed period"),
            (r"\bnet\s+worth\b", "a prescribed financial threshold"),
        ]
        for pattern, what in forbidden:
            m = re.search(pattern, blob, re.I)
            self.assertIsNone(
                m, f"{what} appears in the library — the Rules were not verifiable, "
                   f"so no prescribed detail may be stated: ...{blob[max(0, (m.start() if m else 0) - 60):(m.end() if m else 0) + 40]}...")

    def test_act_citations_do_not_pin_definition_clause_letters(self):
        """Section 2 is cited by defined term, never by clause letter.

        The definitions are lettered alphabetically, which makes the letter a
        function of the whole list rather than of the term — and the list could
        not be checked against the bare Act here. A wrong letter is the same
        defect as a wrong rule number, so it is held to the same standard: cite
        s.2 and name the term.
        """
        for c in LIB.controls.values():
            self.assertIsNone(re.search(r"s\.\s*2\s*\([a-z]\)", c.act_ref),
                              f"{c.id} pins a definition clause letter: {c.act_ref}")
            if re.search(r"s\.\s*2\b", c.act_ref):
                self.assertIn('"', c.act_ref,
                              f"{c.id} cites s.2 without naming the defined term")

    def test_act_citations_name_sections_that_exist(self):
        """The Act runs to 44 sections and a Schedule."""
        for c in LIB.controls.values():
            for num in re.findall(r"s\.\s*(\d+)", c.act_ref):
                self.assertTrue(1 <= int(num) <= 44,
                                f"{c.id} cites s.{num}, outside the Act")

    def test_penalty_entries_cite_the_provision_they_punish(self):
        """Each Schedule entry a control invokes must name an Act provision."""
        entries = {e["entry"]: e for e in LIB.citations["penalty_schedule"]["entries"]}
        for c in LIB.controls.values():
            if c.penalty_entry is None:
                continue
            entry = entries[c.penalty_entry]
            self.assertTrue(str(entry.get("provision") or "").strip(),
                            f"entry {c.penalty_entry} punishes no stated provision")
            # The residuary entry catches every other provision, so only the
            # specific entries are held to overlapping with the control's own
            # citation.
            if entry["provision"] != "residuary":
                head = entry["provision"].split("(")[0]        # "s.8" from "s.8(5)"
                self.assertIn(head, c.act_ref,
                              f"{c.id} claims entry {c.penalty_entry} ({entry['provision']}) "
                              f"but cites {c.act_ref}")


class TestScoping(unittest.TestCase):
    #: Flags that widen the audit. An unanswered question here must leave the
    #: obligation in scope, so that silence never shrinks the engagement.
    WIDENING = ("processes_children_data", "processes_employee_data",
                "has_pre_commencement_data")

    def test_defaults_are_conservative(self):
        """Silence widens scope; every exemption must be claimed explicitly."""
        empty = CompanyProfile.blank()
        for name, (default, _desc) in PROFILE_FLAGS.items():
            self.assertEqual(empty.flag(name), default, name)
            if name in self.WIDENING:
                self.assertTrue(default, f"{name} should default in-scope")
            else:
                self.assertFalse(default, f"{name} should require a positive claim")

    def test_sdf_controls_are_excluded_when_not_designated(self):
        not_sdf = _profile(is_sdf=False)
        ids = {c.id for c in LIB.in_scope(not_sdf)}
        self.assertFalse(any(i.startswith("SDF-") for i in ids))

    def test_sdf_controls_are_included_when_designated(self):
        sdf = _profile(is_sdf=True)
        ids = {c.id for c in LIB.in_scope(sdf)}
        self.assertTrue({"SDF-01", "SDF-02"}.issubset(ids))

    def test_children_controls_follow_the_children_flag(self):
        with_kids = {c.id for c in LIB.in_scope(_profile(processes_children_data=True))}
        without = {c.id for c in LIB.in_scope(_profile(processes_children_data=False))}
        self.assertIn("CHD-03", with_kids)
        self.assertNotIn("CHD-03", without)

    def test_scope_reason_is_recorded_for_excluded_controls(self):
        not_sdf = _profile(is_sdf=False)
        ctrl = LIB.controls["SDF-01"]
        self.assertFalse(ctrl.applies_to(not_sdf))
        self.assertTrue(ctrl.scope_reason(not_sdf).strip())

    def test_scoping_is_reproducible_from_the_flags_alone(self):
        p1 = _profile(is_sdf=True, processes_children_data=True)
        p2 = _profile(is_sdf=True, processes_children_data=True)
        self.assertEqual({c.id for c in LIB.in_scope(p1)},
                         {c.id for c in LIB.in_scope(p2)})


class TestScoringArithmetic(unittest.TestCase):
    def setUp(self):
        self.profile = _profile(is_sdf=True, processes_children_data=True,
                                uses_consent_manager=True)
        self.assessment = Assessment.blank_for(LIB, self.profile)

    def _in_scope(self, domain_code):
        return [c for c in LIB.in_scope(self.profile) if c.domain == domain_code]

    def _score_all(self, value):
        for c in LIB.in_scope(self.profile):
            self.assessment.response(c.id).score = value
        return score_assessment(LIB, self.assessment)

    def test_all_fours_is_one_hundred_percent_and_no_gaps(self):
        r = self._score_all(4)
        self.assertEqual(r.overall_maturity_pct, 100.0)
        self.assertEqual(len(r.gaps), 0)
        self.assertEqual(r.penalty_exposure()["aggregate_max_inr_crore"], 0)

    def test_all_zeroes_is_zero_percent_and_every_control_is_a_gap(self):
        r = self._score_all(0)
        self.assertEqual(r.overall_maturity_pct, 0.0)
        self.assertEqual(len(r.gaps), r.in_scope_count)

    def test_uniform_score_gives_that_score_as_a_percentage(self):
        for value in (0, 1, 2, 3, 4):
            r = self._score_all(value)
            self.assertAlmostEqual(r.overall_maturity_pct, value / 4 * 100, places=1)

    def test_not_applicable_leaves_the_denominator_rather_than_scoring_zero(self):
        self._score_all(4)
        target = self._in_scope("SEC")[0]
        self.assessment.response(target.id).score = NOT_APPLICABLE
        r = score_assessment(LIB, self.assessment)
        self.assertEqual(r.overall_maturity_pct, 100.0,
                         "NA must not drag maturity down")
        self.assertEqual(len(r.gaps), 0, "NA must not manufacture a gap")
        self.assertEqual(r.na_count, 1)

    def test_unassessed_controls_are_a_coverage_limitation_not_a_gap(self):
        self._score_all(4)
        target = self._in_scope("SEC")[0]
        self.assessment.response(target.id).score = None
        r = score_assessment(LIB, self.assessment)
        self.assertEqual(len(r.gaps), 0)
        self.assertLess(r.coverage_pct, 100.0)
        self.assertEqual([x.control.id for x in r.unassessed], [target.id])

    def test_domain_maturity_uses_control_weights(self):
        self._score_all(4)
        dom = LIB.ordered_domains()[0]
        controls = self._in_scope(dom.code)
        heavy = max(controls, key=lambda c: c.weight)
        self.assessment.response(heavy.id).score = 0
        r = score_assessment(LIB, self.assessment)
        got = r.domain_by_code(dom.code).maturity_pct
        den = sum(4 * c.weight for c in controls)
        num = den - 4 * heavy.weight
        self.assertAlmostEqual(got, round(num / den * 100, 1), places=1)

    def test_overall_is_domain_weighted_not_a_flat_control_average(self):
        """A many-control domain must not outvote a heavier, smaller one."""
        self._score_all(4)
        for c in self._in_scope("SEC"):           # 9 controls, domain weight 1.4
            self.assessment.response(c.id).score = 0
        r_sec = score_assessment(LIB, self.assessment)

        self._score_all(4)
        for c in self._in_scope("CMG"):           # 3 controls, domain weight 0.8
            self.assessment.response(c.id).score = 0
        r_cmg = score_assessment(LIB, self.assessment)

        # Zeroing a whole domain removes exactly that domain's share of the
        # weighted total, whatever the control count inside it.
        total_w = sum(d.weight for d in LIB.ordered_domains()
                      if r_sec.domain_by_code(d.code).maturity_pct is not None)
        for results, code in ((r_sec, "SEC"), (r_cmg, "CMG")):
            share = LIB.domains[code].weight
            self.assertAlmostEqual(results.overall_maturity_pct,
                                   round((total_w - share) / total_w * 100, 1),
                                   places=1, msg=code)
        # And the heavier domain must move the needle further despite CMG
        # having a third as many controls.
        self.assertLess(r_sec.overall_maturity_pct, r_cmg.overall_maturity_pct)

    def test_out_of_scope_controls_are_excluded_from_scoring_entirely(self):
        no_sdf = _profile(is_sdf=False, processes_children_data=True)
        a = Assessment.blank_for(LIB, no_sdf)
        for c in LIB.in_scope(no_sdf):
            a.response(c.id).score = 4
        r = score_assessment(LIB, a)
        self.assertEqual(r.overall_maturity_pct, 100.0)
        self.assertIsNone(r.domain_by_code("SDF").maturity_pct)
        self.assertGreater(r.out_of_scope_count, 0)

    def test_a_score_recorded_for_an_out_of_scope_control_is_ignored(self):
        no_sdf = _profile(is_sdf=False)
        a = Assessment.blank_for(LIB, no_sdf)
        a.response("SDF-01").score = 0
        r = score_assessment(LIB, a)
        sdf = [x for x in r.all_results if x.control.id == "SDF-01"][0]
        self.assertIsNone(sdf.score)
        self.assertFalse(sdf.is_gap)


class TestRiskAndPriority(unittest.TestCase):
    def test_risk_is_severity_rank_times_shortfall(self):
        p = _profile(is_sdf=True, processes_children_data=True)
        a = Assessment.blank_for(LIB, p)
        for c in LIB.in_scope(p):
            a.response(c.id).score = 1
        r = score_assessment(LIB, a)
        for res in r.scored:
            self.assertEqual(res.risk, res.control.severity_rank * 3)

    def test_risk_is_bounded_by_zero_and_sixteen(self):
        p = _profile(is_sdf=True, processes_children_data=True)
        for value in (0, 4):
            a = Assessment.blank_for(LIB, p)
            for c in LIB.in_scope(p):
                a.response(c.id).score = value
            for res in score_assessment(LIB, a).scored:
                self.assertGreaterEqual(res.risk, 0)
                self.assertLessEqual(res.risk, 16)

    def test_a_full_score_carries_no_risk(self):
        p = _profile()
        a = Assessment.blank_for(LIB, p)
        for c in LIB.in_scope(p):
            a.response(c.id).score = 4
        for res in score_assessment(LIB, a).scored:
            self.assertEqual(res.risk, 0)

    def test_risk_bands_partition_the_range(self):
        expected = ([("None", 0)] + [("Low", i) for i in (1, 2, 3)]
                    + [("Medium", i) for i in (4, 5, 6, 7)]
                    + [("High", i) for i in (8, 9, 10, 11)]
                    + [("Critical", i) for i in range(12, 17)])
        for name, value in expected:
            self.assertEqual(risk_band(value), name, value)

    def test_priority_thresholds_and_windows(self):
        self.assertEqual(priority_for(16)[:2], ("P1", 30))
        self.assertEqual(priority_for(12)[:2], ("P1", 30))
        self.assertEqual(priority_for(11)[:2], ("P2", 60))
        self.assertEqual(priority_for(8)[:2], ("P2", 60))
        self.assertEqual(priority_for(7)[:2], ("P3", 90))
        self.assertEqual(priority_for(4)[:2], ("P3", 90))
        self.assertEqual(priority_for(3)[:2], ("P4", 180))
        self.assertEqual(priority_for(0)[:2], ("P4", 180))

    def test_gap_threshold_is_the_substantially_compliant_anchor(self):
        p = _profile()
        for value in range(5):
            a = Assessment.blank_for(LIB, p)
            for c in LIB.in_scope(p):
                a.response(c.id).score = value
            gaps = len(score_assessment(LIB, a).gaps)
            if value < SUBSTANTIALLY_COMPLIANT:
                self.assertGreater(gaps, 0, value)
            else:
                self.assertEqual(gaps, 0, value)

    def test_maturity_bands_cover_zero_to_one_hundred(self):
        for pct in (0, 1, 20, 21, 40, 41, 60, 61, 80, 81, 100):
            name, desc = maturity_band(float(pct))
            self.assertTrue(name and desc, pct)


class TestGapRegister(unittest.TestCase):
    def setUp(self):
        self.assessment = Assessment.load(SAMPLE)
        self.results = score_assessment(LIB, self.assessment)
        self.gaps = build_gap_register(self.results)

    def test_one_row_per_gap(self):
        self.assertEqual(len(self.gaps), len(self.results.gaps))

    def test_refs_are_sequential_and_unique(self):
        self.assertEqual([g.ref for g in self.gaps],
                         [f"G{i:03d}" for i in range(1, len(self.gaps) + 1)])

    def test_ordered_by_descending_residual_risk(self):
        risks = [g.risk for g in self.gaps]
        self.assertEqual(risks, sorted(risks, reverse=True))

    def test_priority_matches_the_risk_band(self):
        for g in self.gaps:
            self.assertEqual(g.priority, priority_for(g.risk)[0], g.ref)

    def test_target_state_is_the_controls_own_level_four_descriptor(self):
        for g in self.gaps:
            self.assertEqual(g.target_state, LIB.controls[g.control_id].scoring["4"], g.ref)

    def test_due_dates_derive_from_the_priority_window(self):
        from datetime import date as _date, timedelta
        start = _date.fromisoformat(self.assessment.assessment_date)
        for g in self.gaps:
            if self.results.assessment.responses[g.control_id].target_date:
                continue
            days = priority_for(g.risk)[1]
            self.assertEqual(g.due_date, (start + timedelta(days=days)).isoformat(), g.ref)

    def test_row_width_matches_the_declared_columns(self):
        from dpdpa_audit.scoring import GAP_COLUMNS
        for g in self.gaps:
            self.assertEqual(len(g.as_row()), len(GAP_COLUMNS), g.ref)


class TestPenaltyExposure(unittest.TestCase):
    def setUp(self):
        self.results = score_assessment(LIB, Assessment.load(SAMPLE))
        self.pe = self.results.penalty_exposure()

    def test_aggregate_is_the_sum_of_the_engaged_entries(self):
        self.assertAlmostEqual(
            self.pe["aggregate_max_inr_crore"],
            sum(e["max_inr_crore"] or 0 for e in self.pe["entries"]), places=3)

    def test_maxima_come_from_the_schedule_unaltered(self):
        schedule = {e["entry"]: e.get("max_penalty_inr_crore")
                    for e in LIB.citations["penalty_schedule"]["entries"]}
        for e in self.pe["entries"]:
            self.assertEqual(e["max_inr_crore"], schedule[e["entry"]], e["entry"])

    def test_only_entries_with_a_failing_control_are_engaged(self):
        gap_ids = {r.control.id for r in self.results.gaps}
        for e in self.pe["entries"]:
            self.assertTrue(set(e["controls"]).issubset(gap_ids), e["entry"])
            self.assertTrue(e["controls"], e["entry"])

    def test_the_ceiling_caveat_is_always_carried(self):
        self.assertIn("ceiling", self.pe["basis"] + self.pe["caveat"])
        self.assertIn("33(2)", self.pe["caveat"])

    def test_a_clean_assessment_engages_nothing(self):
        p = _profile(is_sdf=True, processes_children_data=True)
        a = Assessment.blank_for(LIB, p)
        for c in LIB.in_scope(p):
            a.response(c.id).score = 4
        pe = score_assessment(LIB, a).penalty_exposure()
        self.assertEqual(pe["entries"], [])
        self.assertEqual(pe["aggregate_max_inr_crore"], 0)


class TestAssessmentRoundTrip(unittest.TestCase):
    def test_save_and_load_preserves_everything(self):
        original = Assessment.load(SAMPLE)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "a.json")
            original.save(path)
            again = Assessment.load(path)
        self.assertEqual(again.to_dict(), original.to_dict())

    def test_scores_survive_the_round_trip_exactly(self):
        original = Assessment.load(SAMPLE)
        before = score_assessment(LIB, original)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "a.json")
            original.save(path)
            after = score_assessment(LIB, Assessment.load(path))
        self.assertEqual(after.overall_maturity_pct, before.overall_maturity_pct)
        self.assertEqual(len(after.gaps), len(before.gaps))

    def test_the_file_is_plain_reviewable_json(self):
        with open(SAMPLE, encoding="utf-8") as fh:
            d = json.load(fh)
        self.assertIn("profile", d)
        self.assertIn("responses", d)

    def test_blank_carries_standing_scope_language(self):
        a = Assessment.blank_for(LIB, _profile())
        self.assertTrue(a.limitations)
        self.assertTrue(a.scoping_assumptions)


class TestSampleEngagement(unittest.TestCase):
    """The shipped worked example must stay internally consistent."""

    def setUp(self):
        self.assessment = Assessment.load(SAMPLE)
        self.results = score_assessment(LIB, self.assessment)

    def test_every_in_scope_control_is_assessed(self):
        self.assertEqual(self.results.coverage_pct, 100.0)
        self.assertEqual(self.results.unassessed, [])

    def test_no_response_exists_for_an_out_of_scope_control(self):
        in_scope = {c.id for c in LIB.in_scope(self.assessment.profile)}
        scored = {cid for cid, r in self.assessment.responses.items() if r.assessed}
        self.assertEqual(scored - in_scope, set())

    def test_every_gap_carries_a_finding_an_owner_and_an_evidence_ref(self):
        for r in self.results.gaps:
            self.assertTrue(r.finding.strip(), r.control.id)
            self.assertTrue(r.owner.strip(), r.control.id)
            self.assertTrue(r.evidence_ref.strip(), r.control.id)

    def test_evidence_refs_follow_the_workpaper_convention(self):
        for r in self.results.scored:
            if r.evidence_ref:
                self.assertRegex(r.evidence_ref, r"^WP-\d{3}$", r.control.id)

    def test_the_profile_is_fully_described(self):
        p = self.assessment.profile
        for attr in ("legal_name", "sector", "registered_office", "business_description"):
            self.assertTrue(getattr(p, attr).strip(), attr)
        self.assertTrue(p.systems_in_scope)
        self.assertTrue(p.entities_in_scope)

    def test_the_engagement_is_attributed_and_dated(self):
        a = self.assessment
        self.assertTrue(a.auditor_name.strip())
        self.assertTrue(a.auditor_firm.strip())
        self.assertTrue(a.engagement_ref.strip())
        self.assertRegex(a.assessment_date, r"^\d{4}-\d{2}-\d{2}$")

    def test_the_sample_spans_the_scoring_range(self):
        """A uniformly mediocre sample would demonstrate nothing."""
        scores = {r.score for r in self.results.scored}
        self.assertTrue({0, 2, 4}.issubset(scores), scores)

    def test_maturity_is_uneven_across_domains(self):
        pcts = [d.maturity_pct for d in self.results.domains if d.maturity_pct is not None]
        self.assertGreater(max(pcts) - min(pcts), 30.0)

    def test_scoring_is_deterministic(self):
        a = score_assessment(LIB, Assessment.load(SAMPLE))
        b = score_assessment(LIB, Assessment.load(SAMPLE))
        self.assertEqual(a.overall_maturity_pct, b.overall_maturity_pct)
        self.assertEqual([g.ref for g in build_gap_register(a)],
                         [g.ref for g in build_gap_register(b)])

    def test_status_counts_account_for_every_control(self):
        self.assertEqual(sum(self.results.status_counts().values()), len(LIB.controls))


class TestReportGeneration(unittest.TestCase):
    def setUp(self):
        self.results = score_assessment(LIB, Assessment.load(SAMPLE))

    def test_word_report_is_written_and_non_trivial(self):
        from dpdpa_audit.word_report import build_report
        with tempfile.TemporaryDirectory() as tmp:
            path = build_report(self.results, os.path.join(tmp, "r.docx"))
            self.assertTrue(os.path.exists(path))
            self.assertGreater(os.path.getsize(path), 40_000)

    def test_excel_workbook_is_written_with_the_expected_sheets(self):
        from openpyxl import load_workbook
        from dpdpa_audit.excel_report import MATRIX, build_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = build_workbook(self.results, os.path.join(tmp, "m.xlsx"))
            wb = load_workbook(path)
            for sheet in ("Instructions", "Company Profile", MATRIX, "Dashboard",
                          "Evidence Log", "Citations"):
                self.assertIn(sheet, wb.sheetnames)
            self.assertEqual(wb[MATRIX].max_row, len(LIB.controls) + 1)

    def test_abridgement_never_cuts_mid_word(self):
        from dpdpa_audit.word_report import _abridge
        for r in self.results.scored:
            if not r.finding:
                continue
            for limit in (120, 230, 340, 440):
                short = _abridge(r.finding, limit)
                self.assertTrue(
                    short.endswith((".", "!", "?", "…")),
                    f"{r.control.id} at {limit}: {short[-40:]!r}")
                if not short.endswith("…"):
                    self.assertTrue(" ".join(r.finding.split()).startswith(short))

    def test_abridgement_returns_short_text_untouched(self):
        from dpdpa_audit.word_report import _abridge
        self.assertEqual(_abridge("A short finding.", 200), "A short finding.")

    def test_softbreak_does_not_change_the_visible_key(self):
        from dpdpa_audit.word_report import _softbreak
        for key in LIB.citations["rule_slots"]:
            self.assertEqual(_softbreak(key).replace("​", ""), key)

    def test_id_list_summarises_the_tail_without_cutting_an_id(self):
        from dpdpa_audit.word_report import _id_list
        ids = [f"SEC-{i:02d}" for i in range(1, 10)]
        self.assertEqual(_id_list(ids, 9), ", ".join(ids))
        self.assertTrue(_id_list(ids, 3).endswith("and 6 others"))
        self.assertTrue(_id_list(ids, 8).endswith("and 1 other"))


class TestCli(unittest.TestCase):
    def _run(self, argv):
        """Run the CLI, capturing both streams so tested error paths stay quiet."""
        from dpdpa_audit.cli import main
        import io
        import contextlib
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = main(argv)
        return code, out.getvalue()

    def _expect_exit(self, argv):
        import io
        import contextlib
        from dpdpa_audit.cli import main
        with contextlib.redirect_stdout(io.StringIO()), \
                contextlib.redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as ctx:
                main(argv)
        return ctx.exception

    def test_validate_reports_the_library(self):
        code, out = self._run(["--controls", CONTROLS, "validate"])
        self.assertEqual(code, 0)
        self.assertIn("90", out)

    def test_score_json_is_machine_readable_and_matches_the_engine(self):
        code, out = self._run(["--controls", CONTROLS, "score", SAMPLE, "--json"])
        self.assertEqual(code, 0)
        payload = json.loads(out)
        expected = score_assessment(LIB, Assessment.load(SAMPLE))
        self.assertEqual(payload["overall_maturity_pct"], expected.overall_maturity_pct)
        self.assertEqual(payload["gap_count"], len(expected.gaps))
        self.assertEqual(len(payload["domains"]), len(LIB.domains))

    def test_citations_json_lists_the_unverified_slots(self):
        code, out = self._run(["--controls", CONTROLS, "citations", "--json"])
        self.assertEqual(code, 0)
        self.assertEqual(len(json.loads(out)), len(LIB.unverified_slots()))

    def test_init_blank_writes_a_loadable_assessment(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "e.json")
            code, _ = self._run(["--controls", CONTROLS, "init", "--blank", "--out", path])
            self.assertEqual(code, 0)
            a = Assessment.load(path)
            self.assertEqual(len(score_assessment(LIB, a).gaps), 0)

    def test_init_refuses_to_clobber_without_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "e.json")
            self._run(["--controls", CONTROLS, "init", "--blank", "--out", path])
            self._expect_exit(["--controls", CONTROLS, "init", "--blank", "--out", path])
            code, _ = self._run(["--controls", CONTROLS, "init", "--blank",
                                 "--out", path, "--force"])
            self.assertEqual(code, 0)

    def test_report_writes_both_deliverables(self):
        with tempfile.TemporaryDirectory() as tmp:
            code, _ = self._run(["--controls", CONTROLS, "report", SAMPLE, "--out-dir", tmp])
            self.assertEqual(code, 0)
            names = os.listdir(tmp)
            self.assertTrue(any(n.endswith(".docx") for n in names), names)
            self.assertTrue(any(n.endswith(".xlsx") for n in names), names)

    def test_a_missing_assessment_file_exits_rather_than_tracing_back(self):
        self.assertEqual(
            self._expect_exit(["--controls", CONTROLS, "score", "/nonexistent/x.json"]).code, 2)

    def test_an_unknown_domain_filter_is_rejected(self):
        self._expect_exit(["--controls", CONTROLS, "assess", SAMPLE, "--domain", "ZZZ"])

    def test_an_invalid_library_exits_rather_than_tracing_back(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._expect_exit(["--controls", tmp, "validate"])


class TestWebApp(unittest.TestCase):
    """The workbench carries a second implementation of the scoring engine.

    A port is a second implementation whether or not it is called one, so it is
    held to the first: built from the same library, run over the same
    engagement, and required to agree field for field. The two node harnesses do
    the work; this class makes them part of the one test command, and skips
    rather than passes silently when node is absent.
    """

    @classmethod
    def setUpClass(cls):
        if shutil.which("node") is None:
            raise unittest.SkipTest("node is not available; the workbench checks need it")
        cls.tmp = tempfile.TemporaryDirectory()
        cls.app = os.path.join(cls.tmp.name, "app.html")
        r = subprocess.run(
            [sys.executable, os.path.join(ROOT, "webapp", "build_webapp.py"),
             "--controls", CONTROLS, "--sample", SAMPLE, "--out", cls.app],
            capture_output=True, text=True, cwd=ROOT)
        if r.returncode != 0:
            raise AssertionError("the workbench did not build:\n" + r.stderr)

    @classmethod
    def tearDownClass(cls):
        if hasattr(cls, "tmp"):
            cls.tmp.cleanup()

    def _node(self, script, *args):
        return subprocess.run(["node", os.path.join(ROOT, "tests", script), self.app, *args],
                              capture_output=True, text=True, cwd=ROOT)

    def test_the_app_is_self_contained(self):
        with open(self.app, encoding="utf-8") as fh:
            html = fh.read()
        self.assertGreater(len(html), 200_000, "the library should be inlined, not fetched")
        self.assertNotIn("__DATA__", html)
        self.assertNotIn("__SAMPLE__", html)
        # No network reference of any kind: the file has to open from a USB stick.
        for pattern in ("http://", "https://", "src=", "@import", "fetch("):
            self.assertNotIn(pattern, html, f"the app reaches outside itself: {pattern}")
        self.assertIn("color-scheme: light", html)

    def test_the_javascript_engine_agrees_with_the_python_engine(self):
        r = self._node("verify_webapp.mjs", SAMPLE)
        self.assertEqual(r.returncode, 0, r.stderr)
        js = json.loads(r.stdout)

        results = score_assessment(LIB, Assessment.load(SAMPLE))
        gaps = build_gap_register(results)
        self.assertEqual(js["overall_maturity_pct"], results.overall_maturity_pct)
        self.assertEqual(js["overall_band"], results.overall_band[0])
        self.assertEqual(js["controls_in_scope"], results.in_scope_count)
        self.assertEqual(js["controls_assessed"], results.assessed_count)
        self.assertEqual(js["coverage_pct"], results.coverage_pct)
        self.assertEqual(js["gaps_by_severity"], results.gap_counts_by_severity())
        self.assertEqual(js["status_counts"], results.status_counts())
        self.assertEqual(js["penalty_exposure"], results.penalty_exposure())
        self.assertEqual({d["code"]: d["maturity_pct"] for d in js["domains"]},
                         {d.code: d.maturity_pct for d in results.domains})
        # The register is the report's spine, so it is compared row by row.
        self.assertEqual(js["gap_register"], [g.as_row() for g in gaps])

    def test_the_rendered_markup_is_sound(self):
        r = self._node("smoke_webapp.mjs")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("OK", r.stdout)


class TestNoMutationAcrossRuns(unittest.TestCase):
    def test_scoring_does_not_mutate_the_library(self):
        before = copy.deepcopy({c.id: (c.severity, c.weight, dict(c.scoring))
                                for c in LIB.controls.values()})
        score_assessment(LIB, Assessment.load(SAMPLE))
        after = {c.id: (c.severity, c.weight, dict(c.scoring))
                 for c in LIB.controls.values()}
        self.assertEqual(after, before)

    def test_scoring_does_not_mutate_the_assessment(self):
        a = Assessment.load(SAMPLE)
        before = json.dumps(a.to_dict(), sort_keys=True)
        score_assessment(LIB, a)
        self.assertEqual(json.dumps(a.to_dict(), sort_keys=True), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
